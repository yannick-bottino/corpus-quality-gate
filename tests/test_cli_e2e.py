import os, json
from cqg.cli import run

def test_end_to_end(tmp_path):
    from reportlab.pdfgen import canvas
    corpus = tmp_path / "corpus"; corpus.mkdir()
    for name in ("a.pdf", "b.pdf"):
        c = canvas.Canvas(str(corpus / name))
        to = c.beginText(72, 800)
        for _ in range(30):
            to.textLine("Version v1.0 du 2026-01-01.")
        c.drawText(to); c.showPage(); c.save()
    cfg = tmp_path / "cfg.yaml"
    cfg.write_text("llm:\n  provider: mock\n  temperature: 0\npaths:\n  workdir: ./wd\n"
                   "thresholds:\n  coverage_flag_below: 0.7\nparsing:\n  parser: legacy\n",
                   encoding="utf-8")
    out = tmp_path / "out"
    paths = run(str(corpus), str(cfg), str(out))
    assert os.path.exists(paths["corpus_report"]["xlsx"])
    assert os.path.exists(out / "a.score.json")
    data = json.loads((out / "a.score.json").read_text(encoding="utf-8"))
    assert data["config_hash"] and "level" in data


def test_run_isolates_and_flags_bad_document(tmp_path):
    from reportlab.pdfgen import canvas
    corpus = tmp_path / "corpus"; corpus.mkdir()
    c = canvas.Canvas(str(corpus / "good.pdf")); to = c.beginText(72, 800)
    for _ in range(30):
        to.textLine("Version v1.0 du 2026-01-01.")
    c.drawText(to); c.showPage(); c.save()
    (corpus / "corrupt.pdf").write_bytes(b"%PDF-1.4 fichier corrompu \x00\x01\x02 pas un vrai pdf")
    cfg = tmp_path / "cfg.yaml"
    cfg.write_text("llm:\n  provider: mock\nthresholds:\n  coverage_flag_below: 0.7\n"
                   "parsing:\n  parser: legacy\n", encoding="utf-8")
    out = tmp_path / "out"
    result = run(str(corpus), str(cfg), str(out))
    assert result["n_docs"] == 2
    # parse_document is resilient (never crashes): the corrupt PDF is no longer a
    # processing_error, it is short-circuited with an unreadable flag.
    bad = json.loads((out / "corrupt.score.json").read_text(encoding="utf-8"))
    assert "unreadable" in bad["flags"]
    # the valid document must be present and score normally
    good = json.loads((out / "good.score.json").read_text(encoding="utf-8"))
    assert "level" in good


def test_cli_enrich_writes_enriched_md(tmp_path, monkeypatch):
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import ImageReader
    from PIL import Image
    corpus = tmp_path / "corpus"; corpus.mkdir()
    src = tmp_path / "pic.png"; Image.new("RGB", (80, 60), (0, 0, 200)).save(src)
    pdf = corpus / "doc.pdf"; c = canvas.Canvas(str(pdf))
    c.drawString(72, 700, "Texte du document.")
    c.drawImage(ImageReader(str(src)), 72, 500, width=120, height=90); c.showPage(); c.save()
    cfg = tmp_path / "cfg.yaml"
    cfg.write_text("llm:\n  provider: mock\nparsing:\n  parser: legacy\nenrichment:\n  enabled: true\n",
                   encoding="utf-8")
    from cqg.cli import run
    out = tmp_path / "out"
    run(str(corpus), str(cfg), str(out), enrich=True)
    enriched = out / "doc.enriched.md"
    assert enriched.exists()
    assert "description automatique" in enriched.read_text(encoding="utf-8")


def test_cli_enrich_manual_writes_manifest(tmp_path):
    # Manual mode: run() must flush() the manifest of images to describe, otherwise no
    # trace exists to complete the enrichment in-session. The placeholder stays
    # intact in the .enriched.md as long as the description is not filled in.
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import ImageReader
    from PIL import Image
    corpus = tmp_path / "corpus"; corpus.mkdir()
    src = tmp_path / "pic.png"; Image.new("RGB", (80, 60), (200, 0, 0)).save(src)
    pdf = corpus / "doc.pdf"; c = canvas.Canvas(str(pdf))
    c.drawString(72, 700, "Texte du document.")
    c.drawImage(ImageReader(str(src)), 72, 500, width=120, height=90); c.showPage(); c.save()
    manifest_path = tmp_path / "manifest.json"
    cfg = tmp_path / "cfg.yaml"
    cfg.write_text(
        "llm:\n  provider: mock\nparsing:\n  parser: legacy\nenrichment:\n  enabled: true\n"
        f"  vlm:\n    provider: manual\n    manifest_path: {manifest_path.as_posix()}\n",
        encoding="utf-8")
    from cqg.cli import run
    out = tmp_path / "out"
    result = run(str(corpus), str(cfg), str(out))
    from pathlib import Path as _P
    assert _P(result["image_manifest"]) == manifest_path
    assert manifest_path.exists()
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert len(manifest) >= 1
    assert set(("image_path", "placeholder", "description")) <= set(manifest[0].keys())
    enriched = out / "doc.enriched.md"
    assert enriched.exists()
    assert "[[IMAGE" in enriched.read_text(encoding="utf-8")


def test_cli_enrich_feeds_eval_and_flags_auto_descriptions(tmp_path):
    # Enrich-before-eval: the ENRICHED markdown (injected descriptions) is what gets scored,
    # and the presence of auto-generated descriptions is surfaced as a safeguard flag for
    # human review (anti-fabrication: the score relies in part on unverified content).
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import ImageReader
    from PIL import Image
    corpus = tmp_path / "corpus"; corpus.mkdir()
    src = tmp_path / "pic.png"; Image.new("RGB", (80, 60), (0, 0, 200)).save(src)
    pdf = corpus / "doc.pdf"; c = canvas.Canvas(str(pdf))
    c.drawString(72, 700, "Texte du document.")
    c.drawImage(ImageReader(str(src)), 72, 500, width=120, height=90); c.showPage(); c.save()
    cfg = tmp_path / "cfg.yaml"
    cfg.write_text("llm:\n  provider: mock\nparsing:\n  parser: legacy\nenrichment:\n  enabled: true\n"
                   "  vlm:\n    provider: mock\n", encoding="utf-8")
    from cqg.cli import run
    out = tmp_path / "out"
    run(str(corpus), str(cfg), str(out), enrich=True)
    data = json.loads((out / "doc.score.json").read_text(encoding="utf-8"))
    assert any(f.startswith("auto_descriptions:") for f in data["flags"])


def test_cli_enrich_config_only_with_separate_vlm(tmp_path):
    # Config-only path: enrichment.enabled: true drives enrichment without passing
    # enrich=True, with a vlm client distinct from the judgment llm (enrichment.vlm).
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import ImageReader
    from PIL import Image
    corpus = tmp_path / "corpus"; corpus.mkdir()
    src = tmp_path / "pic.png"; Image.new("RGB", (80, 60), (0, 200, 0)).save(src)
    pdf = corpus / "doc.pdf"; c = canvas.Canvas(str(pdf))
    c.drawString(72, 700, "Texte du document.")
    c.drawImage(ImageReader(str(src)), 72, 500, width=120, height=90); c.showPage(); c.save()
    cfg = tmp_path / "cfg.yaml"
    cfg.write_text(
        "llm:\n  provider: mock\nparsing:\n  parser: legacy\nenrichment:\n  enabled: true\n"
        "  vlm:\n    provider: mock\n",
        encoding="utf-8")
    from cqg.cli import run
    out = tmp_path / "out"
    run(str(corpus), str(cfg), str(out))
    enriched = out / "doc.enriched.md"
    assert enriched.exists()
    assert "description automatique" in enriched.read_text(encoding="utf-8")


def test_run_scores_text_and_office_documents(tmp_path):
    # Score-and-flag across formats: .md, .docx and .pptx are all real scored
    # documents. None of them may land as "unreadable" (a document cqg tried and
    # failed to extract) or as a processing error.
    from docx import Document
    from pptx import Presentation
    corpus = tmp_path / "corpus"; corpus.mkdir()
    (corpus / "note.md").write_text(
        "# Notice\n\nVersion v1.0 du 2026-01-01.\n\n" + "Le contrat couvre l'incendie. " * 30,
        encoding="utf-8")

    d = Document()
    for _ in range(20):
        d.add_paragraph("Le contrat couvre l'incendie et le degat des eaux.")
    d.save(str(corpus / "rapport.docx"))

    prs = Presentation()
    for i in range(3):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = f"Garantie {i + 1}"
        s.placeholders[1].text = "Le contrat couvre l'incendie. " * 10
    prs.save(str(corpus / "deck.pptx"))

    cfg = tmp_path / "cfg.yaml"
    cfg.write_text("llm:\n  provider: mock\nparsing:\n  parser: legacy\n", encoding="utf-8")
    out = tmp_path / "out"
    result = run(str(corpus), str(cfg), str(out))
    assert result["n_docs"] == 3
    assert result["n_errors"] == 0

    for name in ("note", "rapport", "deck"):
        ds = json.loads((out / f"{name}.score.json").read_text(encoding="utf-8"))
        assert ds["criteria"], f"{name} must be scored like any other document"
        assert "unreadable" not in ds["flags"]
        assert not any(f.startswith("unsupported_format") for f in ds["flags"])


def test_run_flags_an_admitted_but_unopenable_format(tmp_path, monkeypatch):
    # The unsupported_format exit path is the honest landing for a format triage
    # admits but the parser cannot open. No shipped extension is in that state
    # today, so the branch is driven through triage directly -- keeping the
    # contract covered rather than letting it silently stop being exercised.
    import cqg.cli as cli
    corpus = tmp_path / "corpus"; corpus.mkdir()
    (corpus / "classeur.xlsx").write_bytes(b"PK\x03\x04 pas ouvert par cqg")
    monkeypatch.setattr(cli, "triage_corpus", lambda folder: [
        {"doc_id": "classeur", "type": "xlsx", "hash": "0" * 64,
         "path": str(corpus / "classeur.xlsx"), "pages": None,
         "category": "unsupported_format"}])
    cfg = tmp_path / "cfg.yaml"
    cfg.write_text("llm:\n  provider: mock\nparsing:\n  parser: legacy\n", encoding="utf-8")
    out = tmp_path / "out"
    result = run(str(corpus), str(cfg), str(out))

    ds = json.loads((out / "classeur.score.json").read_text(encoding="utf-8"))
    assert any(f.startswith("unsupported_format") for f in ds["flags"])
    assert "unreadable" not in ds["flags"]
    assert result["n_errors"] == 0, "an unsupported format is not a processing error"
