"""Non-regression gate of lot 1: splitting parsing from scoring changes no score.

`cqg parse` then `cqg run <parsed_input>` must produce exactly the `*.score.json` a
direct `cqg run <raw>` produces -- enrichment included, since the enrichment moved from
run to parse (Q9). Lot 1 carries no LLM precisely so this comparison is interpretable:
anything that moves, moves because of the split.
"""
import json

import pytest

from cqg.cli import parse_corpus, run


def _corpus(tmp_path):
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    from PIL import Image
    corpus = tmp_path / "raw"; corpus.mkdir()

    c = canvas.Canvas(str(corpus / "texte.pdf"))
    to = c.beginText(72, 800)
    for _ in range(30):
        to.textLine("Version v1.0 du 2026-01-01. Le contrat couvre l'incendie.")
    c.drawText(to); c.showPage(); c.save()

    src = tmp_path / "pic.png"; Image.new("RGB", (80, 60), (0, 0, 200)).save(src)
    c = canvas.Canvas(str(corpus / "illustre.pdf"))
    c.drawString(72, 700, "Notice illustree. Redige par le service qualite.")
    c.drawImage(ImageReader(str(src)), 72, 500, width=120, height=90)
    c.showPage(); c.save()

    (corpus / "note.md").write_text(
        "# Notice\n\nVersion v1.0 du 2026-01-01.\n\n" + "Le contrat couvre l'incendie. " * 30,
        encoding="utf-8")

    (corpus / "corrupt.pdf").write_bytes(b"%PDF-1.4 corrompu \x00\x01 pas un vrai pdf")
    return corpus


def _cfg(tmp_path, enrich):
    cfg = tmp_path / "cfg.yaml"
    body = "llm:\n  provider: mock\nparsing:\n  parser: legacy\n"
    if enrich:
        body += "enrichment:\n  enabled: true\n  vlm:\n    provider: mock\n"
    cfg.write_text(body, encoding="utf-8")
    return str(cfg)


def _scores(out_dir):
    return {p.name: json.loads(p.read_text(encoding="utf-8"))
            for p in sorted(out_dir.glob("*.score.json"))}


@pytest.mark.parametrize("enrich", [False, True], ids=["sans_enrichissement", "enrichi"])
def test_parse_then_run_reproduces_a_direct_run(tmp_path, enrich):
    corpus = _corpus(tmp_path)
    cfg = _cfg(tmp_path, enrich)

    direct_out = tmp_path / "direct"
    run(str(corpus), cfg, str(direct_out), enrich=enrich)

    parsed = tmp_path / "parsed"
    parse_corpus(str(corpus), cfg, str(parsed), enrich=enrich)
    split_out = tmp_path / "split"
    # In the enriched case this run deliberately scores with enrichment OFF: it is
    # `parse` that enriched, and `run` on a parsed_input/ warns and ignores the option.
    # So the equality is carried by the markdown parse wrote -- which is the contract.
    run(str(parsed), cfg, str(split_out))

    direct, split = _scores(direct_out), _scores(split_out)
    assert set(direct) == set(split) and len(direct) == 4
    for name in direct:
        assert split[name] == direct[name], f"{name} a bouge entre run direct et parse+run"


def test_editing_the_prose_does_not_flip_the_structure_criteria_to_na(tmp_path):
    """The consequence the Q11 refinement exists for, asserted rather than assumed.

    Recomputing ALL blocks from an edited markdown would return prose blocks only. The
    document's table would vanish from the inventory, na_decisions would flip the table
    criteria to `na`, and the score would move because a human fixed a typo -- the drift
    Q2 forbids. Keeping the typed blocks from the sidecar is what prevents that.
    """
    from docx import Document
    corpus = tmp_path / "raw"; corpus.mkdir()
    d = Document()
    for _ in range(10):
        d.add_paragraph("Version v1.0 du 2026-01-01. Le contrat couvre l'incendie.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Garantie"; table.cell(0, 1).text = "Plafond"
    table.cell(1, 0).text = "Incendie"; table.cell(1, 1).text = "100 000 EUR"
    d.save(str(corpus / "avec_tableau.docx"))
    cfg = _cfg(tmp_path, enrich=False)

    parsed = tmp_path / "parsed"
    parse_corpus(str(corpus), cfg, str(parsed))
    before_out = tmp_path / "before"
    run(str(parsed), cfg, str(before_out))
    before = _scores(before_out)["avec_tableau.score.json"]

    md_path = parsed / "avec_tableau.md"
    edited = md_path.read_bytes().decode("utf-8").replace("l'incendie", "l'incendie et la foudre")
    md_path.write_bytes(edited.encode("utf-8"))
    after_out = tmp_path / "after"
    run(str(parsed), cfg, str(after_out))
    after = _scores(after_out)["avec_tableau.score.json"]

    assert "manually_edited" in after["flags"]
    na_before = {c["id"] for c in before["criteria"] if c["status"] == "na"}
    na_after = {c["id"] for c in after["criteria"] if c["status"] == "na"}
    assert na_after == na_before, "l'edition de la prose a fait basculer des criteres en na"
    # Teeth: the document really does carry a table, so the table criteria are live.
    # Without this the assertion above would hold trivially on a document with no
    # structure at all.
    assert na_before, "le document temoin doit avoir au moins un critere na"
    assert not ({"3.4", "3.8", "3.9"} & na_before), "les criteres tableau doivent etre actifs"
    assert after["global_pct"] == before["global_pct"]


@pytest.mark.parametrize("enrich", [False, True], ids=["sans_enrichissement", "enrichi"])
def test_the_markdown_scored_is_identical_on_both_paths(tmp_path, enrich):
    # Guards the byte-level round trip directly: discovering a lost CR through a score
    # diff would cost hours, discovering it here costs one line.
    corpus = _corpus(tmp_path)
    cfg = _cfg(tmp_path, enrich)
    from cqg.parse import parse_document
    from cqg.enrich import enrich_document
    from cqg.llm.mock import MockLLM
    from cqg import parse_store

    parsed = tmp_path / "parsed"
    parse_corpus(str(corpus), cfg, str(parsed), enrich=enrich)

    doc = parse_document(str(corpus / "illustre.pdf"), pages=1, parser="legacy")
    expected = doc.markdown
    if enrich:
        expected = enrich_document(doc, str(corpus / "illustre.pdf"), MockLLM(),
                                   str(tmp_path / "img"))
    entry = parse_store.read_parsed_doc(str(parsed / "illustre.md"))
    assert entry.doc.markdown == expected
    assert entry.flags == []
