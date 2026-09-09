from cqg.triage import triage_file

def test_born_digital_pdf(tmp_path):
    from reportlab.pdfgen import canvas
    p = tmp_path / "doc.pdf"
    c = canvas.Canvas(str(p))
    to = c.beginText(72, 800)
    for _ in range(40):
        to.textLine("Lorem ipsum dolor sit amet")
    c.drawText(to); c.showPage(); c.save()
    r = triage_file(str(p))
    assert r["category"] == "born_digital"
    assert r["pages"] == 1
    assert len(r["hash"]) == 64


def _make_docx(path, paragraphs):
    from docx import Document
    d = Document()
    for para in paragraphs:
        d.add_paragraph(para)
    d.save(str(path))


def _make_pptx(path, n_slides):
    from pptx import Presentation
    prs = Presentation()
    for i in range(n_slides):
        prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = f"Slide {i + 1}"
    prs.save(str(path))


def test_text_formats_are_classified_parseable(tmp_path):
    # .txt/.md carry extractable text: they are a parseable category, not
    # lumped with formats cqg cannot open.
    from cqg.triage import triage_file
    for name in ("note.md", "note.txt"):
        p = tmp_path / name
        p.write_text("Du contenu.", encoding="utf-8")
        assert triage_file(str(p))["category"] == "text"


def test_office_formats_are_classified_parseable(tmp_path):
    # .docx/.pptx carry structured text that cqg now extracts: they are a
    # parseable category, not a format that is merely admitted and flagged.
    from cqg.triage import triage_file
    docx_path = tmp_path / "doc.docx"; _make_docx(docx_path, ["Du contenu."])
    assert triage_file(str(docx_path))["category"] == "office"

    pptx_path = tmp_path / "deck.pptx"
    _make_pptx(pptx_path, 3)
    assert triage_file(str(pptx_path))["category"] == "office"


def test_pptx_reports_its_slide_count_and_docx_reports_none(tmp_path):
    # `pages` feeds the per-page density check in parse confidence, so it must
    # only ever carry a real count. A pptx has a true slide count; Word
    # pagination is a renderer artefact that python-docx cannot report, so
    # inventing one would feed that check a meaningless number.
    from cqg.triage import triage_file
    pptx_path = tmp_path / "deck.pptx"
    _make_pptx(pptx_path, 3)
    assert triage_file(str(pptx_path))["pages"] == 3

    docx_path = tmp_path / "doc.docx"; _make_docx(docx_path, ["Du contenu."])
    assert triage_file(str(docx_path))["pages"] is None


def test_corrupt_pptx_stays_in_triage_and_defers_to_the_parser(tmp_path):
    # Triage is an inventory, not an extraction verdict: a slide count that
    # cannot be read degrades to None instead of pre-judging the document.
    # The unreadable verdict is the parser's to make, uniformly for both
    # office formats (triage never opens a .docx).
    from cqg.triage import triage_file
    p = tmp_path / "bad.pptx"
    p.write_bytes(b"PK\x03\x04 pas un vrai fichier office")
    r = triage_file(str(p))
    assert r["category"] == "office"
    assert r["pages"] is None


def test_unknown_format_is_classified_unsupported(tmp_path):
    # Score-and-flag: a format cqg cannot parse must stay distinguishable from a
    # document that failed to extract, so it keeps its own category and flag.
    from cqg.triage import triage_file
    p = tmp_path / "classeur.xlsx"
    p.write_bytes(b"PK\x03\x04 pas ouvert par cqg")
    assert triage_file(str(p))["category"] == "unsupported_format"
