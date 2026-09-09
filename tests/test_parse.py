import os
import pytest


def _make_pdf(path, lines, n=20):
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(path)); to = c.beginText(72, 800)
    for _ in range(n):
        for ln in lines:
            to.textLine(ln)
    c.drawText(to); c.showPage(); c.save()


def test_pages_text_pdfminer_nfkc(tmp_path):
    from cqg.parse import _pages_text_pdfminer, _nfkc
    assert _nfkc("ﬁn") == "fin"  # ligature fi -> fi
    p = tmp_path / "d.pdf"; _make_pdf(p, ["Bonjour le monde."])
    pages = _pages_text_pdfminer(str(p))
    assert isinstance(pages, list) and "Bonjour" in "\n".join(pages)


def test_typing_pdfplumber_detects_image(tmp_path):
    from cqg.parse import _typing_pdfplumber
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import ImageReader
    from PIL import Image
    img = tmp_path / "pic.png"; Image.new("RGB", (60, 40), (200, 0, 0)).save(img)
    p = tmp_path / "d.pdf"; c = canvas.Canvas(str(p))
    c.drawImage(ImageReader(str(img)), 72, 600, width=120, height=80)
    c.drawString(72, 500, "Texte."); c.showPage(); c.save()
    blocks, images_by_page = _typing_pdfplumber(str(p))
    assert any(b.kind == "image" for b in blocks)
    assert 0 in images_by_page and len(images_by_page[0]) >= 1


def test_assemble_inserts_placeholders_in_order():
    from cqg.parse import _assemble
    pages = ["Texte page 1", "Texte page 2"]
    imgs = {0: [{"x0": 0, "top": 10, "x1": 5, "bottom": 20},
                {"x0": 0, "top": 50, "x1": 5, "bottom": 60}]}
    md, refs = _assemble(pages, imgs)
    assert "[[IMAGE:p=1;idx=1]]" in md and "[[IMAGE:p=1;idx=2]]" in md
    assert md.index("[[IMAGE:p=1;idx=1]]") < md.index("[[IMAGE:p=1;idx=2]]")
    assert "Texte page 2" in md
    assert [r.placeholder for r in refs] == ["[[IMAGE:p=1;idx=1]]", "[[IMAGE:p=1;idx=2]]"]
    assert refs[0].page == 1 and refs[0].idx == 1


def test_confidence_signals():
    from cqg.parse import _confidence
    from cqg.models import Block
    full = "phrase complete. " * 200
    good = _confidence(full, [Block(kind="text", text=full, page=1)], pages=1, fallback_used=False)
    assert good >= 0.8
    # secondary fallback -> capped at 0.6
    fb = _confidence(full, [Block(kind="text", text=full, page=1)], pages=1, fallback_used=True)
    assert fb <= 0.6
    # very short text over many pages -> low
    assert _confidence("ok", [Block(kind="text", text="ok", page=1)], pages=50, fallback_used=False) <= 0.4


def test_confidence_penalized_by_cid_failure():
    from cqg.parse import _confidence
    from cqg.models import Block
    # Text half composed of unmapped (cid:NNN) tokens: the confidence must
    # NOT stay high (CG Auto case which came out at 0.947 despite the extraction failure).
    clean = "phrase lisible et complete. " * 100
    cid = "(cid:114)(cid:97)(cid:103)(cid:101) " * 100
    mixed = clean + cid
    blocks = [Block(kind="text", text=mixed, page=1)]
    conf_clean = _confidence(clean, blocks, pages=1, fallback_used=False)
    conf_mixed = _confidence(mixed, blocks, pages=1, fallback_used=False)
    assert conf_clean >= 0.8
    assert conf_mixed < 0.7
    assert conf_mixed < conf_clean


def test_better_extraction_prefers_lower_cid():
    from cqg.parse import _better_extraction
    from cqg.models import Block
    cid_md = "(cid:1)(cid:2)(cid:3) " * 50
    clean_md = "texte parfaitement lisible et complet. " * 20
    primary = (cid_md, [Block(kind="text", text=cid_md, page=1)], [])
    fallback = (clean_md, [Block(kind="text", text=clean_md, page=1)], [])
    chosen, used = _better_extraction(primary, fallback)
    assert chosen[0] == clean_md
    assert used is True


def test_better_extraction_keeps_primary_when_clean():
    from cqg.parse import _better_extraction
    from cqg.models import Block
    clean_md = "texte lisible. " * 20
    worse_md = "(cid:9) " * 40
    primary = (clean_md, [Block(kind="text", text=clean_md, page=1)], [])
    fallback = (worse_md, [Block(kind="text", text=worse_md, page=1)], [])
    chosen, used = _better_extraction(primary, fallback)
    assert chosen[0] == clean_md
    assert used is False


def test_parse_document_full_chain(tmp_path):
    from cqg.parse import parse_document
    p = tmp_path / "d.pdf"; _make_pdf(p, ["Contenu complet de la page."])
    doc = parse_document(str(p), parser="legacy")
    assert "Contenu complet" in doc.markdown
    assert 0.0 <= doc.parse_confidence <= 1.0


def test_parse_document_unreadable_never_raises(tmp_path):
    from cqg.parse import parse_document
    bad = tmp_path / "bad.pdf"; bad.write_bytes(b"%PDF-1.4 broken")
    doc = parse_document(str(bad), parser="legacy")
    assert doc.parse_confidence == 0.0
    assert any(b.kind == "unreadable" for b in doc.blocks)


def test_legacy_parser_does_not_use_docling(tmp_path, monkeypatch):
    # Invariant: parser="legacy" never calls Docling and parses via pdfminer/pdfplumber.
    # (Docling is now the DEFAULT parser; legacy mode remains the lightweight fallback.)
    import cqg.parse as parse

    def _boom(*a, **k):
        raise AssertionError("Docling ne doit pas etre utilise par le parser legacy")

    monkeypatch.setattr(parse, "_docling_extraction", _boom)
    p = tmp_path / "d.pdf"; _make_pdf(p, ["Contenu complet de la page."])
    doc = parse.parse_document(str(p), parser="legacy")
    assert "Contenu complet" in doc.markdown


def test_parser_docling_dispatches_to_docling_extraction(tmp_path, monkeypatch):
    # parser="docling" -> the Docling extraction is indeed called and its result used.
    import cqg.parse as parse
    from cqg.models import Block
    called = {}

    def _fake(path, pages=None, batch_pages=None):
        called["yes"] = True
        return ("MD docling structure", [Block(kind="text", text="MD docling structure", page=1)], [])

    monkeypatch.setattr(parse, "_docling_extraction", _fake)
    p = tmp_path / "d.pdf"; _make_pdf(p, ["x"])
    doc = parse.parse_document(str(p), parser="docling")
    assert called.get("yes") is True
    assert doc.markdown == "MD docling structure"


def test_parser_docling_falls_back_to_default_on_failure(tmp_path, monkeypatch):
    # Robustness (score-and-flag): if Docling fails (import/OOM/error), we fall back to
    # the default chain rather than giving up on the document.
    import cqg.parse as parse

    def _boom(path, pages=None, batch_pages=None):
        raise RuntimeError("docling indisponible")

    monkeypatch.setattr(parse, "_docling_extraction", _boom)
    p = tmp_path / "d.pdf"; _make_pdf(p, ["Contenu de repli."])
    doc = parse.parse_document(str(p), parser="docling")
    assert "Contenu de repli" in doc.markdown  # the default chain took over


def _make_pdf_pages(path, page_lines):
    # Multi-page PDF: one page_lines entry = the lines of one page (showPage per page).
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(path))
    for lines in page_lines:
        to = c.beginText(72, 800)
        for ln in lines:
            to.textLine(ln)
        c.drawText(to); c.showPage()
    c.save()


def test_docling_extraction_batches_one_subprocess_per_page(tmp_path, monkeypatch):
    # Batching: batch_pages=1 -> a FRESH subprocess per page. subprocess.run is
    # mocked (no real docling) and writes a fixed markdown into the output file.
    import cqg.parse as parse
    from pathlib import Path
    calls = {"n": 0}

    class _Proc:
        returncode = 0

    def _fake_run(args, **kwargs):
        calls["n"] += 1
        out = next(a for a in args if str(a).endswith(".md"))
        Path(out).write_text("TEXTE_LOT", encoding="utf-8")
        return _Proc()

    monkeypatch.setattr(parse.subprocess, "run", _fake_run)
    p = tmp_path / "d.pdf"
    _make_pdf_pages(p, [["Page une."], ["Page deux."], ["Page trois."]])
    md, blocks, refs = parse._docling_extraction(str(p), batch_pages=1)
    assert calls["n"] == 3
    assert md.count("TEXTE_LOT") == 3


def test_docling_extraction_per_page_pdfminer_fallback_on_oom(tmp_path, monkeypatch):
    # Per-batch fallback: returncode 137 (OOM) -> each page switches to pdfminer without
    # dropping the document. The PDF has real text, so pdfminer produces content.
    import cqg.parse as parse

    class _Proc:
        returncode = 137

    def _fake_run(args, **kwargs):
        return _Proc()

    monkeypatch.setattr(parse.subprocess, "run", _fake_run)
    p = tmp_path / "d.pdf"; _make_pdf(p, ["Contenu reel de la page."])
    md, blocks, refs = parse._docling_extraction(str(p), batch_pages=1)
    assert md.strip()
    assert "Contenu reel" in md


def test_typing_pdfplumber_detects_table(tmp_path):
    from cqg.parse import _typing_pdfplumber
    from reportlab.pdfgen import canvas
    from reportlab.platypus import Table, TableStyle
    from reportlab.lib import colors
    p = tmp_path / "d.pdf"
    c = canvas.Canvas(str(p))
    data = [["A", "B", "C"], ["1", "2", "3"], ["4", "5", "6"], ["7", "8", "9"]]
    table = Table(data, colWidths=80, rowHeights=24)
    # GRID (all inner and outer borders) + BOX (thick outline) are
    # necessary: pdfplumber.find_tables() detects vector lines, not
    # a simple text layout. Without drawn borders, no table
    # is identified.
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 1.5, colors.black),
        ("BOX", (0, 0), (-1, -1), 2, colors.black),
    ]))
    w, h = table.wrapOn(c, 400, 400)
    table.drawOn(c, 72, 700 - h)
    c.showPage(); c.save()
    blocks, _ = _typing_pdfplumber(str(p))
    assert any(b.kind == "table" for b in blocks)


@pytest.mark.skipif(os.getenv("CQG_TEST_DOCLING") != "1",
                    reason="test docling reel opt-in (CQG_TEST_DOCLING=1), lourd (torch)")
def test_docling_real_extraction(tmp_path):
    # Opt-in: checks that the real Docling worker does extract a structured markdown
    # (table rendered with pipes). Skipped by default to keep the suite fast.
    from cqg.parse import parse_document
    from reportlab.pdfgen import canvas
    from reportlab.platypus import Table, TableStyle
    from reportlab.lib import colors
    p = tmp_path / "d.pdf"
    c = canvas.Canvas(str(p))
    data = [["A", "B", "C"], ["1", "2", "3"], ["4", "5", "6"]]
    table = Table(data, colWidths=80, rowHeights=24)
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 1.5, colors.black),
        ("BOX", (0, 0), (-1, -1), 2, colors.black),
    ]))
    _, h = table.wrapOn(c, 400, 400)
    table.drawOn(c, 72, 700 - h)
    c.showPage(); c.save()
    doc = parse_document(str(p), parser="docling")
    assert doc.markdown.strip()
    assert "|" in doc.markdown


def test_parse_document_takes_no_positional_beyond_path(tmp_path):
    # parse_document previously accepted a `category` argument it never read, and
    # both call sites passed it POSITIONALLY. Dropping the parameter without a
    # keyword-only guard would have silently bound the category string to `pages`.
    # Keyword-only makes that class of positional drift impossible.
    import inspect
    from cqg.parse import parse_document
    params = inspect.signature(parse_document).parameters
    assert "category" not in params
    positional = [n for n, p in params.items()
                  if p.kind in (p.POSITIONAL_ONLY, p.POSITIONAL_OR_KEYWORD)]
    assert positional == ["path"]
    p = tmp_path / "d.pdf"; _make_pdf(p, ["Texte."])
    with pytest.raises(TypeError):
        parse_document(str(p), "born_digital")     # ex-category, now rejected


def test_parse_document_extracts_plain_text_and_markdown(tmp_path):
    # .txt/.md need no PDF machinery: read + NFKC. Downstream stages operate on a
    # markdown string, so a text document traverses the pipeline like any other.
    from cqg.parse import parse_document
    p = tmp_path / "note.md"
    p.write_text("# Titre\n\nUn paragraphe de contenu.\n", encoding="utf-8")
    doc = parse_document(str(p))
    assert "Un paragraphe de contenu." in doc.markdown
    assert doc.parse_confidence > 0
    assert doc.images == []
    assert any(b.kind == "text" and b.text.strip() for b in doc.blocks)


def test_parse_document_text_never_uses_pdf_parsers(tmp_path, monkeypatch):
    # A text file must not reach docling or pdfminer.
    import cqg.parse as parse
    def _boom(*a, **k):
        raise AssertionError("PDF parser invoked on a text document")
    monkeypatch.setattr(parse, "_docling_extraction", _boom)
    monkeypatch.setattr(parse, "_pages_text_pdfminer", _boom)
    monkeypatch.setattr(parse, "_pdfplumber_only", _boom)
    p = tmp_path / "note.txt"; p.write_text("Contenu simple.", encoding="utf-8")
    assert "Contenu simple." in parse.parse_document(str(p), parser="docling").markdown


def test_parse_document_undecodable_text_lowers_confidence(tmp_path):
    # Decoding damage becomes U+FFFD, which cid_failure_fraction already counts:
    # a mis-decoded file is penalised and flagged by the existing mechanism
    # rather than passing as clean.
    from cqg.parse import parse_document
    good = tmp_path / "good.txt"
    good.write_text("Le contrat couvre l'incendie. " * 40, encoding="utf-8")
    bad = tmp_path / "bad.txt"
    bad.write_bytes(("Le contrat couvre l'incendie. " * 40).encode("utf-16"))
    assert parse_document(str(bad)).parse_confidence < parse_document(str(good)).parse_confidence


def _make_docx(path, paragraphs, table=None):
    from docx import Document
    d = Document()
    for para in paragraphs:
        d.add_paragraph(para)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for i, row in enumerate(table):
            for j, cell in enumerate(row):
                t.cell(i, j).text = cell
    d.save(str(path))


def _make_pptx(path, slides):
    # slides: list of (title, body, notes)
    from pptx import Presentation
    prs = Presentation()
    layout = prs.slide_layouts[1]
    for title, body, notes in slides:
        s = prs.slides.add_slide(layout)
        s.shapes.title.text = title
        s.placeholders[1].text = body
        if notes:
            s.notes_slide.notes_text_frame.text = notes
    prs.save(str(path))


def test_parse_document_extracts_docx_text_and_types_tables(tmp_path):
    # A .docx is a real scored document, not an admitted-but-unopened format.
    # Tables must be typed as blocks so the deterministic scorers see them the
    # same way they see a PDF's tables.
    from cqg.parse import parse_document
    p = tmp_path / "rapport.docx"
    _make_docx(p, ["Notice d'information", "Le contrat couvre l'incendie."],
               table=[["Garantie", "Plafond"], ["Incendie", "100 000 EUR"]])
    doc = parse_document(str(p))
    assert "Le contrat couvre l'incendie." in doc.markdown
    assert "Incendie" in doc.markdown, "table cell text must reach the markdown"
    assert any(b.kind == "table" for b in doc.blocks)
    assert doc.parse_confidence > 0.0


def test_parse_document_extracts_pptx_slides_and_speaker_notes(tmp_path):
    # Speaker notes are prose a RAG system will ingest: they belong in the
    # extracted text and therefore in the quality verdict.
    from cqg.parse import parse_document
    p = tmp_path / "deck.pptx"
    _make_pptx(p, [("Titre de la slide", "Corps de la slide.", "Note du presentateur.")])
    doc = parse_document(str(p))
    assert "Titre de la slide" in doc.markdown
    assert "Corps de la slide." in doc.markdown
    assert "Note du presentateur." in doc.markdown


def test_parse_document_office_never_uses_pdf_parsers(tmp_path, monkeypatch):
    # Office formats carry structured text: no PDF machinery must be reached,
    # including the docling subprocess.
    import cqg.parse as parse
    def _boom(*a, **k):
        raise AssertionError("PDF parser invoked on an office document")
    monkeypatch.setattr(parse, "_docling_extraction", _boom)
    monkeypatch.setattr(parse, "_pages_text_pdfminer", _boom)
    monkeypatch.setattr(parse, "_pdfplumber_only", _boom)
    docx_path = tmp_path / "d.docx"; _make_docx(docx_path, ["Contenu Word."])
    pptx_path = tmp_path / "d.pptx"; _make_pptx(pptx_path, [("T", "Contenu slide.", None)])
    assert "Contenu Word." in parse.parse_document(str(docx_path), parser="docling").markdown
    assert "Contenu slide." in parse.parse_document(str(pptx_path), parser="docling").markdown


def test_parse_document_office_emits_no_image_refs(tmp_path):
    # ImageRef carries PDF-point geometry consumed by the enrichment cropper.
    # Office formats have no such geometry, so none is fabricated: images are
    # counted as blocks only.
    from cqg.parse import parse_document
    p = tmp_path / "deck.pptx"
    _make_pptx(p, [("T", "Corps.", None)])
    assert parse_document(str(p)).images == []


def test_parse_document_corrupt_office_never_raises(tmp_path):
    # S14 chain: parse_document returns an unreadable document rather than
    # propagating. A corrupt .docx makes python-docx raise PackageNotFoundError.
    from cqg.parse import parse_document
    for name in ("bad.docx", "bad.pptx"):
        p = tmp_path / name
        p.write_bytes(b"PK\x03\x04 pas un vrai fichier office")
        doc = parse_document(str(p))
        assert doc.markdown == ""
        assert doc.parse_confidence == 0.0
        assert any(b.kind == "unreadable" for b in doc.blocks)
