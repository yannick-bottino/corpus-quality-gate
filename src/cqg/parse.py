import subprocess
import sys
import tempfile
import unicodedata
from pathlib import Path
from .models import ParsedDoc, Block, ImageRef
from .signals import cid_failure_fraction
from .triage import TEXT_EXTS, OFFICE_EXTS


def _nfkc(s: str) -> str:
    # Compatibility normalisation: flattens ligatures (U+FB01 fi, U+FB02 fl) and
    # other compatibility forms. No effect on French accented letters.
    return unicodedata.normalize("NFKC", s or "")


def _pages_text_pdfminer(path: str) -> list[str]:
    from pdfminer.high_level import extract_pages
    from pdfminer.layout import LTTextContainer
    pages: list[str] = []
    for layout in extract_pages(path):
        parts = [el.get_text() for el in layout if isinstance(el, LTTextContainer)]
        pages.append(_nfkc("".join(parts)).strip())
    return pages


def _typing_pdfplumber(path: str) -> tuple[list[Block], dict[int, list[dict]]]:
    import pdfplumber
    blocks: list[Block] = []
    images_by_page: dict[int, list[dict]] = {}
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages):
            tables = page.find_tables() or []
            for _ in tables:
                blocks.append(Block(kind="table", text="", page=i + 1))
            imgs = page.images or []
            page_imgs = []
            for im in imgs:
                blocks.append(Block(kind="image", text="", page=i + 1))
                page_imgs.append({"x0": float(im["x0"]), "top": float(im["top"]),
                                  "x1": float(im["x1"]), "bottom": float(im["bottom"])})
            if page_imgs:
                images_by_page[i] = sorted(page_imgs, key=lambda d: d["top"])
            # pdfplumber keeps the object cache of every page: on a large PDF
            # (hundreds of pages) the accumulation saturates memory (OOM). We free
            # the cache per page, so the footprint stays bounded.
            page.flush_cache()
    return blocks, images_by_page


def _assemble(pages_text: list[str],
              images_by_page: dict[int, list[dict]]) -> tuple[str, list[ImageRef]]:
    out_parts: list[str] = []
    refs: list[ImageRef] = []
    for i, text in enumerate(pages_text):
        if text:
            out_parts.append(text)
        for k, im in enumerate(images_by_page.get(i, []), start=1):
            ph = f"[[IMAGE:p={i + 1};idx={k}]]"
            out_parts.append(ph)
            refs.append(ImageRef(page=i + 1, idx=k, x0=im["x0"], top=im["top"],
                                 x1=im["x1"], bottom=im["bottom"], placeholder=ph))
    return "\n\n".join(p for p in out_parts if p), refs


def _confidence(markdown: str, blocks: list[Block], pages: int | None,
                fallback_used: bool) -> float:
    if not markdown.strip():
        return 0.0
    length_ok = min(1.0, len(markdown) / 500.0)
    non_empty = sum(1 for b in blocks if b.text.strip())
    ratio = (non_empty / len(blocks)) if blocks else 1.0
    conf = 0.5 * length_ok + 0.5 * ratio
    # Content loss detection: an abnormally low extraction per page
    # (e.g. parser silently failing on some pages) caps the confidence.
    if pages and pages > 0 and (len(markdown) / pages) < 200:
        conf = min(conf, 0.4)
    if fallback_used:
        conf = min(conf, 0.6)
    # Font-mapping failure: the text is present in volume (high length_ok) but
    # unreadable ((cid:NNN) tokens, unmapped glyphs). The length/block confidence does
    # not see it (CG Auto case at 0.947). Penalty proportional to the lost share.
    cid_frac = cid_failure_fraction(markdown)
    conf = conf * (1.0 - cid_frac)
    return round(conf, 3)


def _pdfplumber_only(path: str) -> tuple[str, list[Block], list[ImageRef]]:
    # Fallback 1: pdfplumber for the text AND the typing (0 native artefact).
    import pdfplumber
    pages_text, blocks, images_by_page = [], [], {}
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages):
            pages_text.append(_nfkc(page.extract_text() or "").strip())
            for _ in (page.find_tables() or []):
                blocks.append(Block(kind="table", text="", page=i + 1))
            page_imgs = []
            for im in (page.images or []):
                blocks.append(Block(kind="image", text="", page=i + 1))
                page_imgs.append({"x0": float(im["x0"]), "top": float(im["top"]),
                                  "x1": float(im["x1"]), "bottom": float(im["bottom"])})
            if page_imgs:
                images_by_page[i] = sorted(page_imgs, key=lambda d: d["top"])
            # Same as _typing_pdfplumber: per-page cache freed to bound memory.
            page.flush_cache()
    md, refs = _assemble(pages_text, images_by_page)
    for i, t in enumerate(pages_text):
        if t:
            blocks.append(Block(kind="text", text=t, page=i + 1))
    return md, blocks, refs


def _better_extraction(primary, fallback):
    """Picks the least degraded extraction (lowest cid fraction).

    primary/fallback: (markdown, blocks, refs) tuples. Returns (chosen, fallback_used).
    Switches to the fallback only if it strictly reduces the cid fraction.
    """
    if cid_failure_fraction(fallback[0]) < cid_failure_fraction(primary[0]):
        return fallback, True
    return primary, False


# Above this cid fraction threshold, we attempt a pdfplumber re-extraction: pdfminer
# produced text in volume but unreadable (font mapping failed).
_CID_FALLBACK_THRESHOLD = 0.1

# Generous Docling worker timeout (large documents). An overrun raises and switches to legacy.
_DOCLING_TIMEOUT = 1800

# Docling processing in page batches, each batch in a FRESH subprocess: the models
# reload on each batch but RAM is freed when the process exits (a whole-doc convert
# OOMs on a ~1.2 GB baseline). 1 page = RSS peak ~1.1 GB, the only batch on a ~1.2 GB baseline.
_DOCLING_BATCH_PAGES = 1


def _docling_extraction(path: str, pages: int | None = None,
                        batch_pages: int | None = None) -> tuple[str, list[Block], list[ImageRef]]:
    # Docling isolated in a subprocess: an OOM (SIGKILL -9/137) is seen via returncode and
    # switches to legacy instead of crashing the pipeline. Raises RuntimeError on any failure.
    # Page-batch processing (fresh process per batch) to bound RAM. Per-batch fallback
    # to pdfminer if a batch fails (OOM/rc!=0/empty): we keep content without losing the doc.
    if batch_pages is None:
        batch_pages = _DOCLING_BATCH_PAGES
    from pypdf import PdfReader
    n = len(PdfReader(path).pages)
    pm_pages: list[str] | None = None  # pdfminer text computed lazily on the 1st failure
    md_parts: list[str] = []
    for start in range(1, n + 1, batch_pages):
        end = min(start + batch_pages - 1, n)
        with tempfile.NamedTemporaryFile(suffix=".md", delete=False) as tmp:
            tmp_out = tmp.name
        try:
            proc = subprocess.run(
                [sys.executable, "-m", "cqg.docling_worker", path, tmp_out, str(start), str(end)],
                timeout=_DOCLING_TIMEOUT, capture_output=True)
            batch_md = ""
            if proc.returncode == 0:
                batch_md = _nfkc(Path(tmp_out).read_text(encoding="utf-8")).strip()
        finally:
            try:
                Path(tmp_out).unlink()
            except OSError:
                pass
        if proc.returncode != 0 or not batch_md:
            # Failed-batch fallback: pdfminer text of the pages involved (1-based -> 0-based index).
            if pm_pages is None:
                pm_pages = _pages_text_pdfminer(path)
            fb = [_nfkc(pm_pages[i]).strip() for i in range(start - 1, end) if i < len(pm_pages)]
            batch_md = "\n\n".join(p for p in fb if p)
        if batch_md:
            md_parts.append(batch_md)
    md = "\n\n".join(md_parts)
    if not md:
        raise RuntimeError("markdown docling vide")
    # Typing (tables/images) + image coords: reuses pdfplumber, which keeps the
    # image -> VLM enrichment path functional. Text blocks come from the docling
    # markdown (structure), split on blank lines.
    type_blocks, images_by_page = _typing_pdfplumber(path)
    blocks = list(type_blocks)
    for para in md.split("\n\n"):
        if para.strip():
            blocks.append(Block(kind="text", text=para.strip(), page=0))
    # Image placeholders grouped by page, same format as _assemble, appended at the end of md.
    refs: list[ImageRef] = []
    ph_parts: list[str] = []
    for i in sorted(images_by_page):
        for k, im in enumerate(images_by_page[i], start=1):
            ph = f"[[IMAGE:p={i + 1};idx={k}]]"
            ph_parts.append(ph)
            refs.append(ImageRef(page=i + 1, idx=k, x0=im["x0"], top=im["top"],
                                 x1=im["x1"], bottom=im["bottom"], placeholder=ph))
    if ph_parts:
        md = md + "\n\n" + "\n\n".join(ph_parts)
    return md, blocks, refs


def _plain_text_extraction(path: str) -> tuple[str, list[Block], list[ImageRef]]:
    # .txt/.md carry their text directly: read + NFKC, no PDF machinery. Decoded as
    # utf-8-sig (strips a BOM) with errors="replace", so decoding damage becomes
    # U+FFFD -- which cid_failure_fraction already counts, penalising parse_confidence
    # and flagging the document instead of letting mojibake pass as clean text.
    md = _nfkc(Path(path).read_text(encoding="utf-8-sig", errors="replace")).strip()
    blocks = [Block(kind="text", text=para.strip(), page=1)
              for para in md.split("\n\n") if para.strip()]
    return md, blocks, []



def _md_table(rows: list[list[str]]) -> str:
    # Rendered as a pipe table so the cell text lands in the markdown that every
    # downstream stage reads. Dropping it would hide content a RAG system will ingest.
    if not rows:
        return ""
    head, body = rows[0], rows[1:]
    out = ["| " + " | ".join(head) + " |",
           "| " + " | ".join("---" for _ in head) + " |"]
    out += ["| " + " | ".join(r) + " |" for r in body]
    return "\n".join(out)


def _docx_extraction(path: str) -> tuple[str, list[Block], list[ImageRef]]:
    # Word carries its own structure: headings, paragraphs and tables are read
    # straight from the package. The body is walked in document order (paragraphs
    # and tables are separate collections in python-docx) so a table keeps its
    # position in the text rather than being appended at the end.
    from docx import Document
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc = Document(path)
    parts: list[str] = []
    blocks: list[Block] = []
    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            para = Paragraph(child, doc)
            text = _nfkc(para.text).strip()
            if not text:
                continue
            # Heading levels are carried into the markdown: the structure criteria
            # are scored on the markdown, so flattening headings into plain
            # paragraphs would cost a document points it has earned.
            style = (para.style.name or "") if para.style is not None else ""
            if style.startswith("Heading"):
                level = style.split()[-1]
                parts.append(("#" * min(int(level), 6) if level.isdigit() else "#") + " " + text)
            else:
                parts.append(text)
            blocks.append(Block(kind="text", text=text, page=0))
        elif isinstance(child, CT_Tbl):
            rows = [[_nfkc(c.text).strip() for c in r.cells] for r in Table(child, doc).rows]
            parts.append(_md_table(rows))
            blocks.append(Block(kind="table", text="", page=0))
    for _ in range(_docx_image_count(doc)):
        blocks.append(Block(kind="image", text="", page=0))
    return "\n\n".join(x for x in parts if x), blocks, []


def _docx_image_count(doc) -> int:
    # Counted from the package parts so floating images count too, not only the
    # inline shapes. Best-effort: the count only feeds the image criteria.
    try:
        return len(list(doc.part.package.image_parts))
    except Exception:
        return len(doc.inline_shapes)


def _pptx_extraction(path: str) -> tuple[str, list[Block], list[ImageRef]]:
    # A deck's text lives in shapes, one slide at a time. Speaker notes are
    # included: they are prose a RAG system will ingest, so they belong in the
    # quality verdict rather than being silently dropped.
    from pptx import Presentation
    prs = Presentation(path)
    parts: list[str] = []
    blocks: list[Block] = []
    for page, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if shape.has_table:
                rows = [[_nfkc(c.text).strip() for c in r.cells] for r in shape.table.rows]
                parts.append(_md_table(rows))
                blocks.append(Block(kind="table", text="", page=page))
                continue
            if _is_picture(shape):
                blocks.append(Block(kind="image", text="", page=page))
                continue
            if shape.has_text_frame:
                text = _nfkc(shape.text_frame.text).strip()
                if text:
                    parts.append(text)
                    blocks.append(Block(kind="text", text=text, page=page))
        if slide.has_notes_slide:
            notes = _nfkc(slide.notes_slide.notes_text_frame.text).strip()
            if notes:
                parts.append(notes)
                blocks.append(Block(kind="text", text=notes, page=page))
    return "\n\n".join(x for x in parts if x), blocks, []


def _is_picture(shape) -> bool:
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    try:
        return shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    except (AttributeError, ValueError, KeyError):
        # python-pptx raises on shape types it cannot map; a shape it cannot
        # classify is simply not counted as an image.
        return False


def _direct_extraction(path: str) -> tuple[str, list[Block], list[ImageRef]]:
    # Formats whose text is readable without any PDF machinery. None of them
    # produces ImageRef geometry: ImageRef carries PDF-point coordinates that the
    # enrichment cropper reads, and neither Word nor PowerPoint exposes anything
    # comparable -- filling those fields would produce silently wrong crops.
    ext = Path(path).suffix.lower()
    if ext in TEXT_EXTS:
        return _plain_text_extraction(path)
    if ext == ".docx":
        return _docx_extraction(path)
    return _pptx_extraction(path)


def parse_document(path: str, *, pages: int | None = None,
                   parser: str = "docling",
                   docling_batch_pages: int | None = None) -> ParsedDoc:
    # Keyword-only past `path`: the extraction strategy is chosen from `parser` and
    # from runtime failures, never from the triage category (which this function used
    # to accept and never read). Keyword-only prevents a caller silently binding a
    # stray positional argument to `pages`.
    doc_id = Path(path).stem
    fallback_used = False
    if Path(path).suffix.lower() in TEXT_EXTS | OFFICE_EXTS:
        try:
            md, blocks, refs = _direct_extraction(path)
        except Exception:
            # S14 chain: never propagates. A corrupt or password-protected office
            # file lands as unreadable -- distinct from unsupported_format, which
            # means cqg has no parser for the extension at all.
            md, blocks, refs = "", [Block(kind="unreadable", text="", page=0)], []
        return ParsedDoc(doc_id=doc_id, markdown=md, blocks=blocks,
                         parse_confidence=_confidence(md, blocks, pages, fallback_used),
                         images=refs)
    # Docling is the default parser. On any exception (import/OOM/error), we fall back
    # to the legacy chain (pdfminer + pdfplumber): score-and-flag robustness, never a crash.
    if parser == "docling":
        try:
            md, blocks, refs = _docling_extraction(path, pages=pages, batch_pages=docling_batch_pages)
            return ParsedDoc(doc_id=doc_id, markdown=md, blocks=blocks,
                             parse_confidence=_confidence(md, blocks, pages, fallback_used),
                             images=refs)
        except Exception:
            pass
    try:
        pages_text = _pages_text_pdfminer(path)
        type_blocks, images_by_page = _typing_pdfplumber(path)
        md, refs = _assemble(pages_text, images_by_page)
        blocks = list(type_blocks)
        for i, t in enumerate(pages_text):
            if t:
                blocks.append(Block(kind="text", text=t, page=i + 1))
        if not md.strip():
            raise ValueError("texte primaire vide")
        # cid fallback: pdfminer extracted volume but unreadable (cid tokens). We try
        # pdfplumber and keep the least degraded extraction. Never raises (S14 chain).
        if cid_failure_fraction(md) > _CID_FALLBACK_THRESHOLD:
            try:
                alt = _pdfplumber_only(path)
                (md, blocks, refs), fallback_used = _better_extraction(
                    (md, blocks, refs), alt)
            except Exception:
                pass
    except Exception:
        try:
            fallback_used = True
            md, blocks, refs = _pdfplumber_only(path)
            if not md.strip():
                raise ValueError("fallback vide")
        except Exception:
            return ParsedDoc(doc_id=doc_id, markdown="",
                             blocks=[Block(kind="unreadable", text="", page=0)],
                             parse_confidence=0.0, images=[])
    return ParsedDoc(doc_id=doc_id, markdown=md, blocks=blocks,
                     parse_confidence=_confidence(md, blocks, pages, fallback_used),
                     images=refs)
