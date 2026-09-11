"""Génère le template Excel de la grille de critères (Q13 + Q23).

Vocabulaire métier côté Excel, codes techniques (D/H/L) côté YAML compilé.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path("config/criteres/grille_template.xlsx")

BLEU = "1F3864"
BLEU_CLAIR = "D9E2F3"
GRIS = "F2F2F2"
JAUNE = "FFF2CC"

TITRE = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
GRAS = Font(name="Calibri", size=11, bold=True)
NORMAL = Font(name="Calibri", size=11)
ITAL = Font(name="Calibri", size=10, italic=True, color="595959")

FOND_TITRE = PatternFill("solid", fgColor=BLEU)
FOND_EX = PatternFill("solid", fgColor=GRIS)
FOND_NOTE = PatternFill("solid", fgColor=JAUNE)
FOND_SECTION = PatternFill("solid", fgColor=BLEU_CLAIR)

BORD = Border(*[Side(style="thin", color="BFBFBF")] * 4)
HAUT = Alignment(vertical="top", wrap_text=True)
CENTRE = Alignment(horizontal="center", vertical="center", wrap_text=True)

TYPES = ["Automatique", "Analysé par l'IA", "Mixte", "Au choix de l'outil"]
DECLENCHEURS = ["aucun", "images", "tableaux", "liens", "formules"]
OUI_NON = ["oui", "non"]


def _entete(ws, cols, row=1):
    for i, (titre, largeur) in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=titre)
        c.font, c.fill, c.alignment, c.border = TITRE, FOND_TITRE, CENTRE, BORD
        ws.column_dimensions[get_column_letter(i)].width = largeur
    ws.row_dimensions[row].height = 34
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _ligne(ws, row, valeurs, fond=None):
    for i, v in enumerate(valeurs, start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.font, c.alignment, c.border = NORMAL, HAUT, BORD
        if fond:
            c.fill = fond


def feuille_criteres(wb):
    ws = wb.active
    ws.title = "Critères"
    cols = [
        ("id\n(identifiant unique)", 10),
        ("dimension\n(numéro, voir feuille Dimensions)", 12),
        ("libellé\nCe que l'on vérifie, en une phrase", 46),
        ("poids\n(1 = accessoire … 5 = décisif)", 10),
        ("type_souhaité\n(liste déroulante)", 20),
        ("intention\nPourquoi ce test ? Qu'est-ce qu'on cherche à savoir ?", 52),
        ("déclencheur_na\nLe test ne s'applique pas si le document n'a pas de…", 20),
        ("dépend_référentiel_externe\n(oui = nécessite une liste de questions métier)", 18),
        ("actif\n(non = test conservé mais non joué)", 10),
    ]
    _entete(ws, cols)

    exemples = [
        ("1.4", "1", "Le document indique un numéro de version ET une date",
         3, "Automatique",
         "Savoir si le document est à jour. Un document sans version ni date ne peut "
         "pas être arbitré face à une version plus récente.",
         "aucun", "non", "oui"),
        ("2.3", "2", "Les informations chiffrées sont exactes et vérifiables",
         5, "Analysé par l'IA",
         "Une erreur de chiffre reprise par l'IA de recherche devient une erreur "
         "servie au client. C'est le risque le plus coûteux.",
         "aucun", "non", "oui"),
        ("3.4", "3", "Les tableaux sont lisibles et correctement structurés",
         4, "Mixte",
         "Un tableau cassé à l'extraction perd son sens. On mesure d'abord la "
         "structure, puis l'IA juge si le contenu reste compréhensible.",
         "tableaux", "non", "oui"),
        ("8.3", "8", "Le document couvre les cas d'usage typiques des utilisateurs",
         5, "Analysé par l'IA",
         "Vérifier que le document répond aux vraies questions posées. Nécessite "
         "la liste des questions métier de référence.",
         "aucun", "oui", "oui"),
    ]
    for r, ex in enumerate(exemples, start=2):
        _ligne(ws, r, ex, FOND_EX)

    note = ws.cell(row=7, column=1,
                   value="Les 4 lignes ci-dessus sont des EXEMPLES à supprimer. "
                         "Saisissez vos critères à partir de la ligne 8. "
                         "Une ligne = un test. Consultez la feuille « Lexique » "
                         "avant de remplir les colonnes type_souhaité et intention.")
    note.font, note.fill, note.alignment = ITAL, FOND_NOTE, HAUT
    ws.merge_cells(start_row=7, start_column=1, end_row=7, end_column=9)
    ws.row_dimensions[7].height = 30

    for col, valeurs in (("E", TYPES), ("G", DECLENCHEURS), ("H", OUI_NON), ("I", OUI_NON)):
        dv = DataValidation(type="list", formula1='"' + ",".join(valeurs) + '"',
                            allow_blank=True, showDropDown=False)
        dv.error = "Valeur non autorisée. Utilisez la liste déroulante."
        dv.errorTitle = "Valeur invalide"
        ws.add_data_validation(dv)
        dv.add(f"{col}8:{col}400")
    dvp = DataValidation(type="whole", operator="between", formula1=1, formula2=5,
                         allow_blank=True)
    dvp.error = "Le poids doit être un entier de 1 à 5."
    dvp.errorTitle = "Poids invalide"
    ws.add_data_validation(dvp)
    dvp.add("D8:D400")


def feuille_dimensions(wb):
    ws = wb.create_sheet("Dimensions")
    _entete(ws, [
        ("dimension\n(numéro)", 12),
        ("nom de la dimension", 34),
        ("poids de la dimension\n(1 = accessoire … 5 = décisif)", 18),
        ("à quoi elle sert, en clair", 62),
    ])
    lignes = [
        ("1", "Métadonnées", 2,
         "Le document se présente-t-il ? Titre, auteur, version, date."),
        ("2", "Exactitude", 5, "L'information est-elle juste, à jour et vérifiable ?"),
        ("3", "Structure et mise en forme", 4,
         "Titres, tableaux, images : le document survit-il à l'extraction automatique ?"),
        ("4", "Clarté et non-redondance", 4,
         "Le texte est-il lisible, sans doublons ni passages contradictoires ?"),
        ("5", "Complétude", 5, "Le sujet est-il traité en entier, sans trous ?"),
        ("6", "Fiabilité et traçabilité", 5,
         "Peut-on remonter à la source de chaque affirmation ?"),
        ("7", "Accessibilité du langage", 3,
         "Le vocabulaire est-il compréhensible par le lecteur visé ?"),
        ("8", "Exploitabilité par une IA de recherche", 3,
         "Le document répond-il à des questions réelles, de façon actionnable ?"),
    ]
    for r, l in enumerate(lignes, start=2):
        _ligne(ws, r, l, FOND_EX)

    r = len(lignes) + 3
    c = ws.cell(row=r, column=1, value="Échelle de notation (commune à tous les critères)")
    c.font, c.fill, c.alignment = GRAS, FOND_SECTION, HAUT
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    _ligne(ws, r + 1, ["échelle_max", 5, "",
                       "Chaque critère est noté de 1 à cette valeur. 1 = absent ou "
                       "inexploitable, valeur médiane = partiellement satisfaisant, valeur "
                       "maximale = exemplaire. Ne changez cette valeur que si vous acceptez "
                       "que les scores passés deviennent incomparables."])

    r += 3
    c = ws.cell(row=r, column=1,
                value="Les valeurs ci-dessus sont celles de la grille livrée par défaut. "
                      "Modifiez-les, ajoutez ou retirez des dimensions : chaque dimension "
                      "citée dans la feuille « Critères » doit exister ici avec un poids, "
                      "sinon la compilation s'arrête et vous le signale.")
    c.font, c.fill, c.alignment = ITAL, FOND_NOTE, HAUT
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    ws.row_dimensions[r].height = 32


def feuille_lexique(wb):
    ws = wb.create_sheet("Lexique")
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 62
    ws.column_dimensions["C"].width = 58
    ws.sheet_properties.tabColor = "A6A6A6"

    ligne = [1]

    def section(titre):
        r = ligne[0]
        c = ws.cell(row=r, column=1, value=titre)
        c.font, c.fill, c.alignment = TITRE, FOND_TITRE, CENTRE
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        ws.row_dimensions[r].height = 24
        ligne[0] += 1

    def entete(*titres):
        r = ligne[0]
        for i, t in enumerate(titres, start=1):
            c = ws.cell(row=r, column=i, value=t)
            c.font, c.fill, c.alignment, c.border = GRAS, FOND_SECTION, CENTRE, BORD
        ligne[0] += 1

    def ligne_(*vals):
        r = ligne[0]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.font, c.alignment, c.border = NORMAL, HAUT, BORD
        ligne[0] += 1

    def vide():
        ligne[0] += 1

    c = ws.cell(row=1, column=1,
                value="Feuille de référence — NE PAS MODIFIER. "
                      "Elle explique les valeurs autorisées dans la feuille « Critères ».")
    c.font, c.fill = ITAL, FOND_NOTE
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)
    ligne[0] = 3

    section("1. Colonne « type_souhaité » : qui juge ce critère ?")
    entete("Valeur à choisir", "Ce que ça veut dire", "Quand l'utiliser")
    ligne_("Automatique",
           "Une règle mécanique lit le texte et décide. Même document = même note, "
           "à chaque fois. Aucun coût d'IA.",
           "Quand la réponse se voit sans interprétation : une date est là ou elle "
           "n'est pas là.")
    ligne_("Analysé par l'IA",
           "L'IA lit le document et porte un jugement argumenté. Elle cite ce sur quoi "
           "elle se fonde. Si elle ne peut pas conclure, elle le dit plutôt que de deviner.",
           "Quand il faut comprendre le sens : exactitude, clarté, ton, pertinence.")
    ligne_("Mixte",
           "Une mesure automatique est calculée d'abord, puis transmise à l'IA qui "
           "l'interprète dans son jugement.",
           "Quand un chiffre aide mais ne suffit pas : « il y a 12 tableaux » ne dit "
           "pas s'ils sont lisibles.")
    ligne_("Au choix de l'outil",
           "Vous décrivez l'intention, l'outil décide seul de la méthode et vous "
           "explique son choix dans le rapport de compilation.",
           "Quand vous ne savez pas si votre test est mesurable mécaniquement.")
    vide()
    ligne_("À SAVOIR",
           "Si vous demandez « Automatique » pour un test qui ne peut pas être mesuré "
           "mécaniquement, l'outil NE le fera PAS passer en douce. Il le confie à l'IA, "
           "vous l'écrit en clair dans le rapport, et vous demande de valider avant "
           "que la grille puisse servir.", "")
    vide()
    vide()

    section("2. Les mesures automatiques disponibles")
    c = ws.cell(row=ligne[0], column=1,
                value="Un test « Automatique » s'appuie sur l'une de ces cinq mesures. "
                      "Vous n'avez rien à écrire ici : décrivez votre intention en français "
                      "dans la colonne « intention », l'outil choisit la mesure et vous dit "
                      "laquelle il a retenue, avec ses paramètres.")
    c.font = ITAL
    ws.merge_cells(start_row=ligne[0], start_column=1, end_row=ligne[0], end_column=3)
    ws.row_dimensions[ligne[0]].height = 30
    ligne[0] += 1
    entete("Mesure", "Ce qu'elle fait", "Exemple d'intention qui y mène")
    ligne_("Rechercher une mention",
           "Cherche un ou plusieurs mots ou expressions dans le document et note selon "
           "qu'ils sont présents ou non.",
           "« Vérifier qu'un auteur ou un rédacteur est identifié. »")
    ligne_("Exiger deux mentions",
           "Cherche deux choses à la fois et note plus haut quand les deux sont là, "
           "moyennement quand une seule l'est.",
           "« Le document doit porter un numéro de version ET une date. »")
    ligne_("Comparer à un seuil",
           "Calcule une proportion mesurable (taux de répétition, part de caractères "
           "illisibles…) et la compare à des paliers.",
           "« Pénaliser les documents dont plus de 5 % des lignes sont dupliquées. »")
    ligne_("Compter des occurrences",
           "Compte combien de fois quelque chose apparaît et note selon un seuil.",
           "« Considérer qu'un document est en format FAQ s'il contient au moins "
           "cinq questions. »")
    ligne_("Mesurer une proportion",
           "Rapporte le nombre d'objets d'un type (images, tableaux, liens) au volume "
           "du document.",
           "« Un document de 50 pages sans aucun tableau est suspect. »")
    vide()
    vide()

    section("3. Colonne « déclencheur_na » : quand le test ne s'applique pas")
    c = ws.cell(row=ligne[0], column=1,
                value="« N/A » signifie « sans objet ». Un critère sans objet n'est ni "
                      "réussi ni raté : il sort du calcul, et ne pénalise donc pas le "
                      "document. Sans cette colonne, un document sans image serait puni "
                      "pour la qualité de ses images.")
    c.font = ITAL
    ws.merge_cells(start_row=ligne[0], start_column=1, end_row=ligne[0], end_column=3)
    ws.row_dimensions[ligne[0]].height = 30
    ligne[0] += 1
    entete("Valeur à choisir", "Le test est ignoré si le document ne contient…", "")
    ligne_("aucun", "Le test s'applique toujours.", "")
    ligne_("images", "… aucune image ni figure.", "")
    ligne_("tableaux", "… aucun tableau.", "")
    ligne_("liens", "… aucun lien hypertexte.", "")
    ligne_("formules", "… aucune formule ni équation.", "")
    vide()
    vide()

    section("4. Les autres colonnes")
    entete("Colonne", "Ce qu'il faut y mettre", "Conséquence si vous vous trompez")
    ligne_("id",
           "Un identifiant unique et stable, par exemple « 2.3 ». Ne réutilisez jamais "
           "l'id d'un critère supprimé.",
           "Deux critères avec le même id : la compilation s'arrête et vous le dit.")
    ligne_("dimension",
           "Le numéro d'une dimension existant dans la feuille « Dimensions ».",
           "Dimension inconnue ou sans poids : la compilation s'arrête et vous le dit.")
    ligne_("libellé",
           "Ce que l'on vérifie, en une phrase, du point de vue du lecteur du document.",
           "Un libellé vague donne un jugement vague : c'est le texte que l'IA lit.")
    ligne_("poids",
           "De 1 à 5, l'importance de ce critère DANS SA dimension.",
           "Tous les poids égaux : plus aucune priorité ne ressort du rapport.")
    ligne_("intention",
           "Pourquoi ce test existe et ce que vous cherchez à savoir. Écrivez-le en "
           "français courant, comme à un collègue.",
           "Colonne vide : le test fonctionne quand même, mais l'IA juge sur le seul "
           "libellé et sera moins fidèle à ce que vous vouliez.")
    ligne_("dépend_référentiel_externe",
           "« oui » si le test a besoin d'une liste de questions métier de référence "
           "que l'outil n'a pas.",
           "« oui » sans référentiel fourni : le critère est déclaré « non évalué », "
           "jamais deviné.")
    ligne_("actif",
           "« non » met le test en sommeil : il reste écrit dans la grille mais n'est "
           "pas joué.",
           "Aucune : c'est la façon propre de retirer un test sans perdre son historique.")
    vide()
    vide()

    section("5. Ce qui se passe après que vous avez rendu ce fichier")
    entete("Étape", "Ce qui est produit", "Qui regarde")
    ligne_("1. Compilation",
           "L'outil lit votre classeur, choisit une méthode par test, et produit une "
           "copie « grille_compilee.xlsx » contenant un onglet « Rapport ».",
           "Automatique")
    ligne_("2. Lecture du rapport",
           "L'onglet « Rapport » liste, test par test : la mesure retenue ET SES "
           "PARAMÈTRES en clair, ainsi que toute décision prise à votre place (test "
           "confié à l'IA, test désactivé, valeur corrigée).",
           "Équipe DMO + référent métier")
    ligne_("3. Validation",
           "Si l'outil a dû décider à votre place sur au moins un test, la grille est "
           "bloquée tant qu'une personne ne l'a pas validée. Sinon elle est utilisable "
           "immédiatement.",
           "Référent métier")
    ligne_("4. Évaluation",
           "La grille validée sert à noter le corpus. Chaque rapport de notation porte "
           "l'empreinte de la grille utilisée : deux rapports ne sont comparables que "
           "si cette empreinte est identique.",
           "Équipe DMO")


def main():
    wb = Workbook()
    feuille_criteres(wb)
    feuille_dimensions(wb)
    feuille_lexique(wb)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"écrit: {OUT}")


if __name__ == "__main__":
    main()
